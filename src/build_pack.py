"""Builds the DirectSportsGoods FY26 management pack (P1).

DirectSportsGoods is fictional and so is all of its data. Benchmarks and their
sources are listed in docs/assumptions.md.

Run: python src/build_pack.py
"""
import csv
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
PACK = ROOT / "pack" / "DirectSportsGoods_FY26_management_pack.xlsx"
DATA_CSV = ROOT / "data" / "fy26_monthly.csv"

FICTIONAL = "DirectSportsGoods is a fictional company. All figures are made up."
MONTHS = ["Feb-25", "Mar-25", "Apr-25", "May-25", "Jun-25", "Jul-25",
          "Aug-25", "Sep-25", "Oct-25", "Nov-25", "Dec-25", "Jan-26"]
CAL = [2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 12, 1]  # calendar month of each period
CHANNELS = ["Website", "Marketplace"]

COMPANY = "Company assumption (fictional)"
# key: (label, value, number format, source)
A = {
    "revenue":   ("Budget net revenue, FY26", 48_000_000, "£#,##0", COMPANY),
    "web_share": ("Website share of net revenue", 0.70, "0%", COMPANY),
    "aov_web":   ("Average order value, website (gross)", 62, "£0.00", COMPANY),
    "aov_mkt":   ("Average order value, marketplace (gross)", 45, "£0.00", COMPANY),
    "gm":        ("Gross margin % of net revenue", 0.511, "0.0%", "Frasers AR 2026, p.24 (UK Sports segment)"),
    "payroll":   ("Payroll % of net revenue", 0.141, "0.0%", "Frasers AR 2026, p.153 and p.22 (group-wide, a ceiling for online)"),
    "h1":        ("May-Oct share of a May-Apr year's sales", 0.517, "0.0%", "Frasers HY26 announcement p.2, AR 2026 p.24"),
    "inv_days":  ("Inventory days", 172, "0", "Frasers AR 2026, p.120 and p.118"),
    "pay_days":  ("Trade and other payables days", 118, "0", "Frasers AR 2026, p.120 and p.118"),
    "returns":   ("Returns % of gross sales", 0.282, "0.0%", "Debenhams Group (formerly boohoo) AR 2026, p.17 (fashion-heavy, likely a ceiling)"),
    "delivery":  ("Delivery and fulfilment cost per order", 4.20, "£0.00", "Debenhams Group AR 2026: distribution costs p.94 / orders p.17"),
    "mkt_fee":   ("Marketplace fees % of marketplace net revenue", 0.15, "0.0%", "Amazon UK seller pricing, Sports and Outdoors referral fee (sell.amazon.co.uk/pricing, read 4 Oct 2026)"),
    "marketing": ("Marketing budget % of website net revenue", 0.068, "0.0%", "ASOS AR 2025, p.49 (6.8% of revenue; ASOS sells through its own site)"),
    "overheads": ("Other overheads per month", 300_000, "£#,##0", COMPANY),
    "dep":       ("Depreciation per month", 60_000, "£#,##0", COMPANY),
    "capex":     ("Capital expenditure per month", 70_000, "£#,##0", COMPANY),
    "cash0":     ("Opening cash, 1 Feb 2025", 4_000_000, "£#,##0", COMPANY),
    "tax":       ("Corporation tax main rate", 0.25, "0%", "HMRC main rate from 1 April 2023"),
    "growth":    ("Planned sales growth on FY25", 0.10, "0%", COMPANY),
}
V = {k: v[1] for k, v in A.items()}

# Monthly sales shape (calendar month -> weight). The shape is our estimate;
# it is scaled so May-Oct matches the Frasers half-year split.
RAW = {1: 9, 2: 6, 3: 7, 4: 7, 5: 7.5, 6: 8, 7: 8.5, 8: 9, 9: 8, 10: 7.5, 11: 11, 12: 11.5}
_h1 = sum(RAW[m] for m in range(5, 11))
SEASON = {m: w * (V["h1"] / _h1 if 5 <= m <= 10 else (1 - V["h1"]) / (sum(RAW.values()) - _h1))
          for m, w in RAW.items()}
W = [SEASON[m] for m in CAL]

PL_LINES = ["Gross sales", "Returns", "Cost of goods sold", "Delivery",
            "Marketplace fees", "Marketing", "Payroll", "Other overheads", "Depreciation"]


def model(scn, rng):
    """Monthly P&L lines for one scenario: {(channel, line): [12 values]}. Costs are negative."""
    rows = {}
    for ch in CHANNELS:
        share = V["web_share"] if ch == "Website" else 1 - V["web_share"]
        aov = V["aov_web"] if ch == "Website" else V["aov_mkt"]
        cols = {k: [] for k in ["Orders", "Gross sales", "Returns", "Cost of goods sold",
                                "Delivery", "Marketplace fees", "Marketing"]}
        for p in range(12):
            plan_net = V["revenue"] * share * W[p]
            o, a, r, gm, d = plan_net / (aov * (1 - V["returns"])), aov, V["returns"], V["gm"], V["delivery"]
            mk = plan_net * V["marketing"] if ch == "Website" else 0
            if scn != "Budget":
                o *= rng.gauss(1, 0.03)
                a *= rng.gauss(1, 0.01)
                r += rng.gauss(0, 0.005)
                gm += rng.gauss(0, 0.004)
                d *= rng.gauss(1, 0.02)
                mk *= rng.gauss(1, 0.03)
            if scn == "Prior year":
                o /= 1 + V["growth"]
                mk /= 1 + V["growth"]
            if scn == "Actual":
                # Planted stories for the variance commentary.
                if ch == "Website" and CAL[p] in (11, 12):
                    o *= 0.91  # Black Friday and Christmas conversion below plan
                if ch == "Website" and CAL[p] == 11:
                    mk *= 1.25  # extra paid search to chase Black Friday
                if ch == "Marketplace" and p >= 4:
                    o *= 1.18  # wider marketplace range listed from June
                if p >= 7:
                    d *= 1.12  # carrier price rise from September
                if CAL[p] == 1:
                    gm -= 0.04  # January clearance discounting
            gross = o * a
            net = gross * (1 - r)
            cols["Orders"].append(o)
            cols["Gross sales"].append(gross)
            cols["Returns"].append(-gross * r)
            cols["Cost of goods sold"].append(-net * (1 - gm))
            cols["Delivery"].append(-o * d)
            cols["Marketplace fees"].append(-net * V["mkt_fee"] if ch == "Marketplace" else 0)
            cols["Marketing"].append(-mk)
        for k, vals in cols.items():
            rows[(ch, k)] = vals

    scale = 1 / (1 + V["growth"]) if scn == "Prior year" else 1
    noise = (lambda s: rng.gauss(1, s)) if scn != "Budget" else (lambda s: 1)
    pay = V["revenue"] * V["payroll"] * scale
    rows[("Central", "Payroll")] = [-pay * (0.8 / 12 + 0.2 * W[p]) * noise(0.015) for p in range(12)]
    rows[("Central", "Other overheads")] = [-V["overheads"] * noise(0.04) for _ in range(12)]
    rows[("Central", "Depreciation")] = [-V["dep"]] * 12
    return rows


def total(rows, line, p=None):
    vals = [v for (ch, ln), v in rows.items() if ln == line]
    return sum(v[p] for v in vals) if p is not None else sum(sum(v) for v in vals)


def op_profit(rows):
    return sum(total(rows, ln) for ln in PL_LINES)


def build_data():
    rng = random.Random(2026)
    scen = {s: model(s, rng) for s in ["Budget", "Actual", "Prior year"]}
    act = scen["Actual"]
    # Working capital and cash items (actuals only). Period 0 = opening balance at 31 Jan 2025.
    cogs = [-total(act, "Cost of goods sold", p) for p in range(12)]
    # Stock covers the next 12 months' COGS at Frasers' stock days; payables follow the last 12 months'.
    py = [-total(scen["Prior year"], "Cost of goods sold", p) for p in range(12)]
    fwd = cogs + [c * (1 + V["growth"]) for c in cogs]
    back = py + cogs
    cash_rows = {
        ("Central", "Inventory"): [V["inv_days"] / 365 * sum(fwd[p:p + 12]) for p in range(13)],
        ("Central", "Trade payables"): [V["pay_days"] / 365 * sum(back[p:p + 12]) for p in range(13)],
        ("Central", "Capex"): [None] + [-V["capex"]] * 12,
        # Profits above £1.5m mean quarterly instalments (gov.uk): FY25's last two land in Feb-25 and May-25,
        # FY26's first two in Aug-25 and Nov-25.
        ("Central", "Tax paid"): [None] + [-V["tax"] * op_profit(scen["Prior year" if CAL[p] in (2, 5) else "Actual"]) / 4
                                           if CAL[p] in (2, 5, 8, 11) else 0 for p in range(12)],
    }
    out = []
    for s, rows in scen.items():
        for (ch, ln), vals in rows.items():
            out += [(s, ch, ln, p + 1, MONTHS[p], round(v, 2)) for p, v in enumerate(vals)]
    for (ch, ln), vals in cash_rows.items():
        out += [("Actual", ch, ln, p, "Opening" if p == 0 else MONTHS[p - 1], round(v, 2))
                for p, v in enumerate(vals) if v is not None]
    return scen, out


# ---------- workbook ----------

NAVY = "1F3864"
F_TITLE = Font(bold=True, size=14, color=NAVY)
F_NOTE = Font(italic=True, size=9, color="7F7F7F")
F_HEAD = Font(bold=True, color="FFFFFF")
FILL_HEAD = PatternFill("solid", fgColor=NAVY)
FILL_INPUT = PatternFill("solid", fgColor="DDEBF7")
INPUT, LINK = "0000FF", "008000"  # blue = hard-coded input, green = link to another sheet
TOP = Border(top=Side(style="thin"))
GBP = '#,##0,;(#,##0,);"-"'  # pounds shown in £k
PCT = '0.0%;-0.0%;"-"'


def header(ws, title, note=None):
    ws["A1"] = title
    ws["A1"].font = F_TITLE
    ws["A2"] = FICTIONAL
    ws["A2"].font = F_NOTE
    if note:
        ws["A3"] = note
        ws["A3"].font = F_NOTE
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 34


def head_row(ws, r, labels, start=1):
    for i, t in enumerate(labels):
        c = ws.cell(r, start + i, t)
        c.font, c.fill = F_HEAD, FILL_HEAD
        c.alignment = Alignment(horizontal="center" if i else "left")


def sumifs(scn, ch, line, period):
    return f'SUMIFS(d_val,d_scn,{scn},d_ch,{ch},d_line,"{line}",d_per,{period})'


def style_row(ws, r, ncols, fmt, bold=False, border=False):
    for c in range(2, ncols + 1):
        ws.cell(r, c).number_format = fmt
    for c in range(1, ncols + 1):
        if bold:
            ws.cell(r, c).font = Font(bold=True)
        if border:
            ws.cell(r, c).border = TOP


# P&L layout: (label, data line or formula over earlier labels, kind)
PL = [
    ("Gross sales", "Gross sales", "line"),
    ("Returns", "Returns", "line"),
    ("Net revenue", ["Gross sales", "Returns"], "sub"),
    ("Cost of goods sold", "Cost of goods sold", "line"),
    ("Gross profit", ["Net revenue", "Cost of goods sold"], "sub"),
    ("Gross margin %", ("Gross profit", "Net revenue"), "pct"),
    ("Delivery and fulfilment", "Delivery", "line"),
    ("Marketplace fees", "Marketplace fees", "line"),
    ("Marketing", "Marketing", "line"),
    ("Contribution", ["Gross profit", "Delivery and fulfilment", "Marketplace fees", "Marketing"], "sub"),
    ("Contribution %", ("Contribution", "Net revenue"), "pct"),
    ("Payroll", "Payroll", "line"),
    ("Other overheads", "Other overheads", "line"),
    ("EBITDA", ["Contribution", "Payroll", "Other overheads"], "sub"),
    ("EBITDA %", ("EBITDA", "Net revenue"), "pct"),
    ("Depreciation", "Depreciation", "line"),
    ("Operating profit", ["EBITDA", "Depreciation"], "sub"),
]


def write_pl_rows(ws, r0, cols, cell_formula):
    """Writes the PL layout from row r0. cols: list of column numbers.
    cell_formula(line, col) returns the formula body for a data line."""
    where = {}
    for i, (label, src, kind) in enumerate(PL):
        r = r0 + i
        where[label] = r
        ws.cell(r, 1, label)
        for c in cols:
            L = ws.cell(r, c).column_letter
            if kind == "line":
                f = cell_formula(src, c)
            elif kind == "sub":
                f = "+".join(f"{L}{where[s]}" for s in src)
            else:
                f = f'IF({L}{where[src[1]]}=0,0,{L}{where[src[0]]}/{L}{where[src[1]]})'
            ws.cell(r, c, "=" + f)
        style_row(ws, r, max(cols), PCT if kind == "pct" else GBP,
                  bold=kind == "sub", border=kind == "sub")
        if kind == "pct":
            for c in range(1, max(cols) + 1):
                ws.cell(r, c).font = Font(italic=True, color="595959")
    return where


def build_workbook(scen, data, commentary):
    wb = Workbook()

    # Cover
    ws = wb.active
    ws.title = "Cover"
    header(ws, "DirectSportsGoods: FY26 management pack")
    ws.column_dimensions["A"].width = 100
    lines = [
        "Year ended 31 January 2026. Online sports goods retailer selling through its own website and third-party marketplaces.",
        "",
        "Contents",
        "  Commentary: the three biggest variances against budget, and what to do about them",
        "  P&L: monthly profit and loss. Pick the channel and scenario in cells C3 and C4.",
        "  Budget vs Actual: full-year variances in £ and %",
        "  KPIs: like-for-like sales, gross margin, payroll ratio, average order value",
        "  Cash: opening cash, operating cash flow, closing cash",
        "  Assumptions: every driver and where it comes from",
        "  Data: the monthly source data every page reads from",
        "",
        "Figures are in £k unless stated. Costs show in brackets. Positive variances are favourable.",
        "Benchmarks come from the published accounts of Frasers Group, Debenhams Group and ASOS, and Amazon UK's seller fees.",
    ]
    for i, t in enumerate(lines, start=4):
        ws.cell(i, 1, t).font = Font(bold=t == "Contents")

    # Commentary
    ws = wb.create_sheet("Commentary")
    header(ws, "Commentary: FY26 against budget")
    ws.column_dimensions["A"].width = 110
    r = 4
    for title, paras in commentary:
        ws.cell(r, 1, title).font = Font(bold=True, size=12, color=NAVY)
        r += 1
        for para in paras:
            c = ws.cell(r, 1, para)
            c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.row_dimensions[r].height = 15 * (len(para) // 118 + 1)
            r += 1
        r += 1

    # P&L
    ws = wb.create_sheet("P&L")
    header(ws, "Monthly P&L (£k)")
    ws["A3"], ws["C3"] = "Channel", "Total"
    ws["A4"], ws["C4"] = "Scenario", "Actual"
    for ref in ("C3", "C4"):
        ws[ref].fill = FILL_INPUT
        ws[ref].font = Font(bold=True, color=INPUT)
    dv_ch = DataValidation(type="list", formula1='"Total,Website,Marketplace"')
    dv_sc = DataValidation(type="list", formula1='"Actual,Budget,Prior year"')
    ws.add_data_validation(dv_ch)
    ws.add_data_validation(dv_sc)
    dv_ch.add("C3")
    dv_sc.add("C4")
    ch = 'IF($C$3="Total","*",$C$3)'
    head_row(ws, 6, [""] + MONTHS + ["FY26"])
    cols = list(range(2, 14))
    where = write_pl_rows(ws, 7, cols, lambda line, c: sumifs("$C$4", ch, line, c - 1))
    for label, r in where.items():  # FY column
        kind = next(k for l, s, k in PL if l == label)
        ws.cell(r, 14, f"=SUM(B{r}:M{r})" if kind != "pct" else ws.cell(r, 13).value.replace("M", "N"))
        ws.cell(r, 14).number_format = PCT if kind == "pct" else GBP
        ws.cell(r, 14).font = Font(bold=True, italic=kind == "pct")
    ws.cell(7 + len(PL) + 1, 1, "Payroll, overheads and depreciation are central costs, so they only appear in the Total view.").font = F_NOTE
    for c in range(2, 15):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 10
    ws.freeze_panes = "B7"

    # Budget vs Actual
    ws = wb.create_sheet("Budget vs Actual")
    header(ws, "Budget vs Actual, FY26 (£k)", "Positive variance = favourable. Costs are negative, so spending less than budget shows as positive.")
    head_row(ws, 5, ["", "Actual", "Budget", "Variance £", "Variance %"])
    where = write_pl_rows(ws, 6, [2, 3], lambda line, c: sumifs('"Actual"' if c == 2 else '"Budget"', '"*"', line, '">=1"'))
    for label, r in where.items():
        kind = next(k for l, s, k in PL if l == label)
        if kind == "pct":
            ws.cell(r, 4, f"=B{r}-C{r}").number_format = '+0.0%;-0.0%;"-"'
            ws.cell(r, 4).font = Font(italic=True, color="595959")
        else:
            ws.cell(r, 4, f"=B{r}-C{r}").number_format = '+#,##0,;-#,##0,;"-"'
            ws.cell(r, 5, f'=IF(C{r}=0,"",D{r}/ABS(C{r}))').number_format = '+0.0%;-0.0%;"-"'
            for c in (4, 5):
                ws.cell(r, c).font = Font(bold=kind == "sub")
    r = 6 + len(PL) + 2
    head_row(ws, r, ["By channel", "Actual", "Budget", "Variance £", "Variance %"])
    for chn in CHANNELS:
        for line_label, lines in (("net revenue", ["Gross sales", "Returns"]),
                                  ("contribution", ["Gross sales", "Returns", "Cost of goods sold", "Delivery", "Marketplace fees", "Marketing"])):
            r += 1
            ws.cell(r, 1, f"{chn} {line_label}")
            arr = "{" + ",".join(f'"{x}"' for x in lines) + "}"
            for c, s in ((2, "Actual"), (3, "Budget")):
                ws.cell(r, c, f'=SUM(SUMIFS(d_val,d_scn,"{s}",d_ch,"{chn}",d_line,{arr}))').number_format = GBP
            ws.cell(r, 4, f"=B{r}-C{r}").number_format = '+#,##0,;-#,##0,;"-"'
            ws.cell(r, 5, f'=IF(C{r}=0,"",D{r}/ABS(C{r}))').number_format = '+0.0%;-0.0%;"-"'
    for c, w in zip("BCDE", (12, 12, 12, 12)):
        ws.column_dimensions[c].width = w

    # KPIs
    ws = wb.create_sheet("KPIs")
    header(ws, "KPIs, FY26", "Like-for-like = same channels, this year vs last year. Average order value is before returns.")
    head_row(ws, 5, [""] + MONTHS + ["FY26", "Budget"])
    net = lambda s, per: f'({sumifs(s, chr(34) + "*" + chr(34), "Gross sales", per)}+{sumifs(s, chr(34) + "*" + chr(34), "Returns", per)})'
    one = lambda s, line, per: sumifs(s, '"*"', line, per)
    kpis = [
        ("Net revenue (£k)", lambda s, per: net(s, per), GBP),
        ("Net revenue last year (£k)", lambda s, per: net('"Prior year"', per), GBP),
        ("Like-for-like sales growth", lambda s, per: f'{net(s, per)}/{net(chr(34) + "Prior year" + chr(34), per)}-1', PCT),
        ("Gross margin %", lambda s, per: f'1+{one(s, "Cost of goods sold", per)}/{net(s, per)}', PCT),
        ("Payroll % of net revenue", lambda s, per: f'-{one(s, "Payroll", per)}/{net(s, per)}', PCT),
        ("Orders (000s)", lambda s, per: one(s, "Orders", per), '#,##0,'),
        ("Average order value (£)", lambda s, per: f'{one(s, "Gross sales", per)}/{one(s, "Orders", per)}', "£0.00"),
        ("Returns rate", lambda s, per: f'-{one(s, "Returns", per)}/{one(s, "Gross sales", per)}', PCT),
    ]
    for i, (label, fn, fmt) in enumerate(kpis):
        r = 6 + i
        ws.cell(r, 1, label)
        for p in range(1, 13):
            ws.cell(r, 1 + p, "=" + fn('"Actual"', p)).number_format = fmt
        ws.cell(r, 14, "=" + fn('"Actual"', '">=1"')).number_format = fmt
        ws.cell(r, 15, "=" + fn('"Budget"', '">=1"') if label != "Net revenue last year (£k)" else "").number_format = fmt
        ws.cell(r, 14).font = Font(bold=True)
    for c in range(2, 16):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 10
    ws.freeze_panes = "B6"

    # Cash
    ws = wb.create_sheet("Cash")
    header(ws, "Cash, FY26 (£k)", "Operating cash flow = EBITDA plus working capital movements. Stock and payables are set by Frasers' stock and creditor days.")
    head_row(ws, 5, [""] + MONTHS + ["FY26"])
    ebitda_lines = "{" + ",".join(f'"{x}"' for x in PL_LINES if x != "Depreciation") + "}"
    rows = ["Opening cash", "EBITDA", "(Increase)/decrease in inventory", "Increase/(decrease) in payables",
            "Operating cash flow", "Capital expenditure", "Corporation tax paid", "Net cash flow", "Closing cash"]
    R = {name: 6 + i for i, name in enumerate(rows)}
    for i, name in enumerate(rows):
        ws.cell(R[name], 1, name)
    for p in range(1, 13):
        c = 1 + p
        L, prev = ws.cell(1, c).column_letter, ws.cell(1, c - 1).column_letter
        bal = lambda line, per: f'SUMIFS(d_val,d_scn,"Actual",d_line,"{line}",d_per,{per})'
        f = {
            "Opening cash": "=Assumptions!$B$" + str(5 + list(A).index("cash0")) if p == 1 else f"={prev}{R['Closing cash']}",
            "EBITDA": f'=SUM(SUMIFS(d_val,d_scn,"Actual",d_line,{ebitda_lines},d_per,{p}))',
            "(Increase)/decrease in inventory": f"={bal('Inventory', p - 1)}-{bal('Inventory', p)}",
            "Increase/(decrease) in payables": f"={bal('Trade payables', p)}-{bal('Trade payables', p - 1)}",
            "Operating cash flow": f"=SUM({L}{R['EBITDA']}:{L}{R['Increase/(decrease) in payables']})",
            "Capital expenditure": "=" + bal("Capex", p),
            "Corporation tax paid": "=" + bal("Tax paid", p),
            "Net cash flow": f"={L}{R['Operating cash flow']}+{L}{R['Capital expenditure']}+{L}{R['Corporation tax paid']}",
            "Closing cash": f"={L}{R['Opening cash']}+{L}{R['Net cash flow']}",
        }
        for name, formula in f.items():
            ws.cell(R[name], c, formula)
    for name, r in R.items():
        if name == "Opening cash":
            ws.cell(r, 14, f"=B{r}")
        elif name == "Closing cash":
            ws.cell(r, 14, f"=M{r}")
        else:
            ws.cell(r, 14, f"=SUM(B{r}:M{r})")
        style_row(ws, r, 14, GBP, bold=name in ("Operating cash flow", "Closing cash", "Opening cash"),
                  border=name in ("Operating cash flow", "Net cash flow", "Closing cash"))
    ws.cell(R["Opening cash"], 2).font = Font(bold=True, color=LINK)
    for c in range(2, 15):
        ws.column_dimensions[ws.cell(1, c).column_letter].width = 10
    ws.freeze_panes = "B6"

    # Assumptions
    ws = wb.create_sheet("Assumptions")
    header(ws, "Assumptions", "Blue = hard-coded input. Full sources in docs/assumptions.md.")
    head_row(ws, 4, ["Driver", "Value", "Source"])
    for i, (label, value, fmt, src) in enumerate(A.values()):
        r = 5 + i
        ws.cell(r, 1, label)
        ws.cell(r, 2, value).number_format = fmt
        ws.cell(r, 2).font = Font(color=INPUT)
        ws.cell(r, 3, src)
    r = 6 + len(A)
    ws.cell(r, 1, "Monthly sales shape (share of FY budget)").font = Font(bold=True)
    ws.cell(r, 3, "Our estimate, scaled so May-Oct matches the Frasers half-year split")
    for p in range(12):
        ws.cell(r + 1 + p, 1, MONTHS[p])
        ws.cell(r + 1 + p, 2, round(W[p], 4)).number_format = "0.0%"
        ws.cell(r + 1 + p, 2).font = Font(color=INPUT)
    ws.column_dimensions["A"].width = 46
    ws.column_dimensions["B"].width = 14
    ws.column_dimensions["C"].width = 70

    # Data
    ws = wb.create_sheet("Data")
    head_row(ws, 1, ["Scenario", "Channel", "Line", "Period", "Month", "Value"])
    for row in data:
        ws.append(row)
    n = len(data) + 1
    for name, col in (("d_scn", "A"), ("d_ch", "B"), ("d_line", "C"), ("d_per", "D"), ("d_val", "F")):
        wb.defined_names[name] = DefinedName(name, attr_text=f"Data!${col}$2:${col}${n}")
    ws.column_dimensions["C"].width = 20
    ws.freeze_panes = "A2"

    PACK.parent.mkdir(exist_ok=True)
    wb.save(PACK)


def commentary(scen):
    """Variance commentary. Numbers come from the model so the text never drifts from the pack."""
    a, b = scen["Actual"], scen["Budget"]
    k = lambda x: f"£{abs(x) / 1000:,.0f}k"
    m = lambda x: f"£{x / 1e6:.2f}m"

    def net(r, ch=None, ps=range(12)):
        chs = [ch] if ch else CHANNELS
        return sum(r[(c, "Gross sales")][p] + r[(c, "Returns")][p] for c in chs for p in ps)

    def contrib(r, ch):
        return sum(sum(r[(ch, ln)]) for ln in PL_LINES[:6])

    def orders(r, ps=range(12)):
        return sum(r[(c, "Orders")][p] for c in CHANNELS for p in ps)

    na, nb, pa, pb = net(a), net(b), op_profit(a), op_profit(b)
    gm_a = 1 + total(a, "Cost of goods sold") / na
    mkt_var, mkt_late = net(a, "Marketplace") - net(b, "Marketplace"), net(a, "Marketplace", range(4, 12)) - net(b, "Marketplace", range(4, 12))
    mkt_cm = contrib(a, "Marketplace") / net(a, "Marketplace")
    web_cm = contrib(a, "Website") / net(a, "Website")
    web_var, web_peak = net(a, "Website") - net(b, "Website"), net(a, "Website", [9, 10]) - net(b, "Website", [9, 10])
    nov_mktg = a[("Website", "Marketing")][9] - b[("Website", "Marketing")][9]
    oa, ob = orders(a), orders(b)
    da, db = -total(a, "Delivery"), -total(b, "Delivery")
    vol, rate = (oa - ob) * db / ob, (da / oa - db / ob) * oa
    late = -sum(a[(c, "Delivery")][p] for c in CHANNELS for p in range(7, 12)) / orders(a, range(7, 12))
    jan_hit = net(a, ps=[11]) * 0.04

    return [
        ("Summary", [
            f"Net revenue was {m(na)}, {k(na - nb)} ({(na - nb) / nb:+.1%}) ahead of budget. Operating profit was {m(pa)}, "
            f"{k(pa - pb)} ({(pa - pb) / pb:+.1%}) behind. Sales grew in the channel that earns least, and costs per order rose. "
            f"Gross margin finished at {gm_a:.1%} against {V['gm']:.1%} budgeted, mostly from January clearance discounting (about {k(jan_hit)}).",
        ]),
        (f"1. Marketplace sales {k(mkt_var)} ahead of budget ({mkt_var / net(b, 'Marketplace'):+.1%})", [
            f"What happened: the beat is all from June onwards ({k(mkt_late)} ahead from June to January), after the wider range went live on marketplaces.",
            f"Why it matters: every £1 of marketplace sales earned {mkt_cm:.0%} contribution against {web_cm:.0%} on the website, "
            f"because of marketplace fees. The extra sales added {k(contrib(a, 'Marketplace') - contrib(b, 'Marketplace'))} of contribution.",
            "What to do: keep the range on marketplaces, but use them to sell slower lines and clear stock, and steer repeat customers "
            "to the website. Renegotiate fee tiers now that volumes are higher.",
        ]),
        (f"2. Website sales {k(web_var)} behind budget ({web_var / net(b, 'Website'):+.1%})", [
            f"What happened: {k(web_peak)} of the shortfall was in November and December. Black Friday and Christmas orders came in "
            f"below plan, even after an extra {k(nov_mktg)} of paid search in November.",
            f"Why it matters: this is the higher-margin channel. Website contribution was {k(contrib(a, 'Website') - contrib(b, 'Website'))} "
            "behind budget, more than the whole profit gap. The marketplace beat only partly offset it.",
            "What to do: review the peak trading plan before next November: stock depth on top sellers, site speed and checkout "
            "conversion, and whether extra search spend paid back. Set a spend cap per order for peak weeks.",
        ]),
        (f"3. Delivery costs {k(da - db)} over budget", [
            f"What happened: {k(vol)} came from handling more orders than planned, which is fine. The other {k(rate)} is rate: "
            f"the carrier's price rise in September took cost per order from £{db / ob:.2f} to £{late:.2f} for September to January.",
            "Why it matters: the rise is permanent, so next year's delivery bill is higher unless something changes.",
            "What to do: retender the carrier contract, test a higher free-delivery threshold, and pass part of the increase on "
            "in the delivery charge. Put the new rate in next year's budget.",
        ]),
    ]


def main():
    scen, data = build_data()
    # Self-check: budget hits the revenue target and the seasonality honours the Frasers split.
    b_net = total(scen["Budget"], "Gross sales") + total(scen["Budget"], "Returns")
    assert abs(b_net - V["revenue"]) < 1, b_net
    assert abs(sum(SEASON[m] for m in range(5, 11)) - V["h1"]) < 1e-9
    with open(DATA_CSV, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["scenario", "channel", "line", "period", "month", "value"])
        w.writerows(data)
    notes = commentary(scen)
    build_workbook(scen, data, notes)
    print(f"Wrote {PACK.relative_to(ROOT)} and {DATA_CSV.relative_to(ROOT)}")
    for title, paras in notes:
        print(f"\n{title}\n" + "\n".join(paras))
    for s, rows in scen.items():
        print(f"  {s:10s} net revenue {(total(rows, 'Gross sales') + total(rows, 'Returns')) / 1e6:6.2f}m  op profit {op_profit(rows) / 1e6:5.2f}m")


if __name__ == "__main__":
    main()
