"""Builds the DirectSportsGoods rolling 18-month forecast and CFO memo (P3).

DirectSportsGoods is fictional and so is all of its data. The forecast starts from
FY26 actuals (the P1 pack, which the P2 close reproduces to the penny) and runs from
February 2026 to July 2027 under base, downside and upside scenarios. The workbook
holds live formulas; the memo quotes the same model run in Python.

Run: python src/forecast.py
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter as col
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import build_pack as bp
from build_pack import FILL_INPUT, GBP, INPUT, PCT, head_row, header, style_row

OUT = bp.ROOT / "pack" / "DirectSportsGoods_forecast.xlsx"
MEMO = bp.ROOT / "docs" / "cfo_memo.md"
N = 18
FM = [date(2026 + (1 + i) // 12, (1 + i) % 12 + 1, 1) for i in range(N)]  # Feb-26 to Jul-27
CH = bp.CHANNELS
SCN = ["Base", "Downside", "Upside"]
ONCOST = 0.12  # employer NI and pension on top of pay, as in the P2 payroll data

# key, label, base, downside, upside. Scenario judgements, not benchmarks.
DRIVERS = [
    ("web_growth", "Website orders vs same month last year", 0.05, -0.03, 0.10),
    ("mkt_growth", "Marketplace orders vs same month last year", 0.15, 0.05, 0.25),
    ("aov", "Average order value vs last year", 0.02, 0.00, 0.03),
    ("returns_pp", "Returns rate change (points)", 0.00, 0.01, -0.01),
    ("gm_pp", "Gross margin change (points)", 0.00, -0.01, 0.005),
    ("carrier", "Carrier rate change from today", 0.00, 0.05, 0.00),
    ("productivity", "Warehouse orders per hour change", 0.00, -0.05, 0.05),
    ("pay_rise", "Salaried pay rise each April", 0.03, 0.04, 0.03),
    ("nlw_2027", "National Living Wage rise, April 2027", 0.03, 0.04, 0.03),
    ("overheads", "Overhead inflation", 0.03, 0.04, 0.02),
]
# Options for the CFO memo. Quotes and response rates are fictional company assumptions.
OPTIONS = {
    "quote": (4.45, "Option A: second carrier's quoted rate per order", "£0.00"),
    "switch_cost": (150_000, "Option A: one-off cost to switch carrier", "£#,##0"),
    "charge_exvat": (round(1.99 / 1.2, 4), "Option B: delivery charge per order, ex VAT (£1.99 inc VAT)", "£0.00"),
    "affected": (0.35, "Option B: share of website orders below the free-delivery threshold", "0%"),
    "loss": (0.08, "Option B: share of those orders lost to the charge", "0%"),
    "price_rise": (0.02, "Option C: price rise, both channels", "0.0%"),
    "volume_loss": (0.025, "Option C: orders lost to the price rise", "0.0%"),
}


def calibrate():
    """FY26 actuals (positive magnitudes) and the fixed inputs derived from them."""
    scen, data = bp.build_data()
    a = scen["Actual"]
    A = {k: [abs(x) for x in v] for k, v in a.items()}
    late = range(7, 12)  # Sep-25 to Jan-26, after the carrier price rise
    rate = sum(A[(c, "Delivery")][p] for c in CH for p in late) / sum(A[(c, "Orders")][p] for c in CH for p in late)
    web_net = sum(A[("Website", "Gross sales")]) - sum(A[("Website", "Returns")])
    nlw_fy26 = [11.44 if p < 2 else 12.21 for p in range(12)]  # FY26 runs Feb-25 to Jan-26
    hours = sum(0.45 * pay / (n * (1 + ONCOST)) for pay, n in zip(A[("Central", "Payroll")], nlw_fy26))
    orders = sum(sum(A[(c, "Orders")]) for c in CH)
    bal = {(ln, per): v for s, ch, ln, per, _, v in data if s == "Actual" and ln in ("Inventory", "Trade payables", "Capex", "Tax paid")}
    ebitda = sum(sum(v) for (ch, ln), v in a.items() if ln in bp.PL_LINES and ln != "Depreciation")
    cash = (bp.V["cash0"] + ebitda + bal[("Inventory", 0)] - bal[("Inventory", 12)] + bal[("Trade payables", 12)]
            - bal[("Trade payables", 0)] + sum(v for (ln, p), v in bal.items() if ln in ("Capex", "Tax paid")))
    fixed = {
        "delivery_rate": (round(rate, 4), "Delivery cost per order today (Sep-25 to Jan-26 average)", "£0.00", "Derived from FY26 actuals"),
        "mkt_ratio": (round(sum(A[("Website", "Marketing")]) / web_net, 4), "Marketing % of website net revenue", "0.0%", "FY26 actual"),
        "fee_rate": (bp.V["mkt_fee"], "Marketplace fees % of net revenue", "0.0%", "Amazon UK seller fees, Sports and Outdoors"),
        "orders_per_hour": (round(orders / hours, 3), "Warehouse orders per labour hour", "0.00", "Derived: FY26 orders / warehouse hours (45% of payroll at NLW)"),
        "oncost": (ONCOST, "Employer NI and pension on pay", "0%", "Company assumption, as in the P2 payroll data"),
        "nlw_2025": (12.21, "National Living Wage from April 2025", "£0.00", "gov.uk national minimum wage rates"),
        "nlw_2026": (12.71, "National Living Wage from April 2026", "£0.00", "gov.uk national minimum wage rates"),
        "salaried_month": (round(0.55 * sum(A[("Central", "Payroll")]) / 12, 2), "Salaried payroll per month, FY26 average", "£#,##0", "FY26 actual (55% of payroll)"),
        "inv_days": (bp.V["inv_days"], "Inventory days", "0", "Frasers AR 2026"),
        "pay_days": (bp.V["pay_days"], "Trade and other payables days", "0", "Frasers AR 2026"),
        "capex_m": (bp.V["capex"], "Capital expenditure per month", "£#,##0", "Company assumption"),
        "dep_m": (bp.V["dep"], "Depreciation per month", "£#,##0", "Company assumption"),
        "tax_rate": (bp.V["tax"], "Corporation tax rate", "0%", "HMRC main rate"),
        "fy26_profit": (round(bp.op_profit(a), 2), "FY26 operating profit (last two tax instalments due Feb-26 and May-26)", "£#,##0", "P1 pack"),
        "large_threshold": (1_500_000, "Profit above which tax is paid in quarterly instalments", "£#,##0", "gov.uk, Corporation Tax: paying in instalments"),
        "cash_open": (round(cash, 2), "Cash at 31 January 2026", "£#,##0", "P1 pack, Cash page"),
    }
    return A, fixed


def forecast(A, X, d):
    """One scenario, month by month. Costs negative, as in P1. X: fixed inputs, d: scenario drivers."""
    F = {}
    for ch, g in (("Website", d["web_growth"]), ("Marketplace", d["mkt_growth"])):
        r = {k: [] for k in ["orders", "aov", "gross", "rr", "returns", "net", "gm", "cogs", "delivery", "marketing", "fees", "contribution"]}
        for c in range(N):
            m = c % 12
            ly = lambda line: A[(ch, line)][m]
            r["orders"].append((ly("Orders") if c < 12 else r["orders"][c - 12]) * (1 + g))
            r["aov"].append((ly("Gross sales") / ly("Orders") if c < 12 else r["aov"][c - 12]) * (1 + d["aov"]))
            r["gross"].append(r["orders"][c] * r["aov"][c])
            r["rr"].append(ly("Returns") / ly("Gross sales") + d["returns_pp"])
            r["returns"].append(-r["gross"][c] * r["rr"][c])
            r["net"].append(r["gross"][c] + r["returns"][c])
            r["gm"].append(1 - ly("Cost of goods sold") / (ly("Gross sales") - ly("Returns")) + d["gm_pp"])
            r["cogs"].append(-r["net"][c] * (1 - r["gm"][c]))
            r["delivery"].append(-r["orders"][c] * X["delivery_rate"] * (1 + d["carrier"]))
            r["marketing"].append(-r["net"][c] * X["mkt_ratio"] if ch == "Website" else 0)
            r["fees"].append(-r["net"][c] * X["fee_rate"] if ch == "Marketplace" else 0)
            r["contribution"].append(sum(r[k][c] for k in ("net", "cogs", "delivery", "marketing", "fees")))
        F[ch] = r
    tot = [F["Website"]["orders"][c] + F["Marketplace"]["orders"][c] for c in range(N)]
    nlw = [X["nlw_2025"] if dt < date(2026, 4, 1) else X["nlw_2026"] if dt < date(2027, 4, 1) else X["nlw_2026"] * (1 + d["nlw_2027"]) for dt in FM]
    hours = [o / (X["orders_per_hour"] * (1 + d["productivity"])) for o in tot]
    wh = [-h * n * (1 + X["oncost"]) for h, n in zip(hours, nlw)]
    sal = [-X["salaried_month"] * (1 + d["pay_rise"]) ** ((dt >= date(2026, 4, 1)) + (dt >= date(2027, 4, 1))) for dt in FM]
    oh = [-A[("Central", "Other overheads")][c % 12] * (1 + d["overheads"]) ** (1 + (c >= 12)) for c in range(N)]
    ebitda = [F["Website"]["contribution"][c] + F["Marketplace"]["contribution"][c] + wh[c] + sal[c] + oh[c] for c in range(N)]
    op = [e - X["dep_m"] for e in ebitda]

    cogs_all = [sum(A[(ch, "Cost of goods sold")][p] for ch in CH) for p in range(12)] + \
               [-(F["Website"]["cogs"][c] + F["Marketplace"]["cogs"][c]) for c in range(N)]
    trail = [sum(cogs_all[k:k + 12]) for k in range(N + 1)]  # k=0 is the opening position
    inv = [X["inv_days"] / 365 * t for t in trail]
    pay = [X["pay_days"] / 365 * t for t in trail]
    # Quarterly instalments: FY26's last two in Feb-26 and May-26, FY27's in Aug-26, Nov-26, Feb-27 and May-27.
    # A year with profit under £1.5m pays in one go nine months after year end, outside this forecast.
    fy27 = sum(op[:12])
    fy27_inst = -X["tax_rate"] * fy27 / 4 if fy27 > X["large_threshold"] else 0
    tax = [-X["tax_rate"] * X["fy26_profit"] / 4 if c in (0, 3) else fy27_inst if c in (6, 9, 12, 15) else 0 for c in range(N)]
    cash, close = X["cash_open"], []
    for c in range(N):
        cash += ebitda[c] + inv[c] - inv[c + 1] + pay[c + 1] - pay[c] - X["capex_m"] + tax[c]
        close.append(cash)
    return {**{f"{ch}:{k}": v for ch in CH for k, v in F[ch].items()},
            "orders": tot, "nlw": nlw, "hours": hours, "warehouse": wh, "salaried": sal, "overheads": oh,
            "ebitda": ebitda, "op": op, "cash": close}


def options(F, X, O, d):
    """Monthly cash impact of each option against doing nothing (contribution only, before tax and working capital)."""
    cpo = {ch: [(F[f"{ch}:net"][c] + F[f"{ch}:cogs"][c] + F[f"{ch}:delivery"][c] + F[f"{ch}:fees"][c]) / F[f"{ch}:orders"][c]
                for c in range(N)] for ch in CH}
    rate = X["delivery_rate"] * (1 + d["carrier"])
    a = [-O["switch_cost"] if c == 1 else (rate - O["quote"]) * F["orders"][c] if c >= 3 else 0 for c in range(N)]
    b = []
    for c in range(N):
        hit = F["Website:orders"][c] * O["affected"]
        b.append(hit * (1 - O["loss"]) * O["charge_exvat"] - hit * O["loss"] * cpo["Website"][c] if c >= 2 else 0)
    cc = []
    for c in range(N):
        gain = 0
        for ch in CH:
            o, net = F[f"{ch}:orders"][c], F[f"{ch}:net"][c]
            lift = O["price_rise"] * net / o * (1 - (X["fee_rate"] if ch == "Marketplace" else 0))
            gain += o * (1 - O["volume_loss"]) * (cpo[ch][c] + lift) - o * cpo[ch][c]
        cc.append(gain if c >= 2 else 0)
    fy = range(12)
    web_cpo = sum(cpo["Website"][c] * F["Website:orders"][c] for c in fy) / sum(F["Website:orders"][c] for c in fy)
    lift = O["price_rise"] * sum(F["Website:net"][c] + F["Marketplace:net"][c] * (1 - X["fee_rate"]) for c in fy)
    contrib = sum(cpo[ch][c] * F[f"{ch}:orders"][c] for ch in CH for c in fy)
    payback = next((c for c in range(2, N) if sum(a[:c + 1]) >= 0), None)  # first month the switch cost is recovered
    return {"A": a, "B": b, "C": cc, "cpo_web": web_cpo, "be_B": O["charge_exvat"] / (O["charge_exvat"] + web_cpo),
            "be_C": lift / (contrib + lift), "payback_A": payback}


# ---------- workbook ----------

def build_workbook(A, X, runs):
    wb = Workbook()
    ws = wb.active
    ws.title = "Cover"
    header(ws, "DirectSportsGoods: rolling 18-month forecast")
    ws.column_dimensions["A"].width = 100
    for i, t in enumerate([
        "February 2026 to July 2027, built from FY26 actuals. Pick a scenario on the Inputs sheet (cell C3).",
        "",
        "Inputs: scenario drivers, fixed inputs and option assumptions",
        "Forecast: monthly P&L and cash, all live formulas",
        "Options: cash impact of the three options in the CFO memo, for the selected scenario",
        "Summary: all three scenarios side by side (values, rebuilt by src/forecast.py)",
        "Actuals: FY26 monthly actuals the forecast starts from",
        "",
        "Figures are in £k unless stated. Costs show in brackets. Blue cells are inputs.",
    ], start=4):
        ws.cell(i, 1, t)

    # Inputs
    ws = wb.create_sheet("Inputs")
    header(ws, "Inputs", "Scenario drivers are judgements, not benchmarks. The Live column feeds the forecast.")
    ws["A3"], ws["C3"] = "Scenario", "Base"
    ws["C3"].fill, ws["C3"].font = FILL_INPUT, Font(bold=True, color=INPUT)
    dv = DataValidation(type="list", formula1='"Base,Downside,Upside"')
    ws.add_data_validation(dv)
    dv.add("C3")
    head_row(ws, 5, ["Scenario driver", "Base", "Downside", "Upside", "Live"])
    for i, (key, label, *vals) in enumerate(DRIVERS):
        r = 6 + i
        ws.cell(r, 1, label)
        for j, v in enumerate(vals):
            ws.cell(r, 2 + j, v).number_format = '+0.0%;-0.0%;0.0%'
            ws.cell(r, 2 + j).font = Font(color=INPUT)
        ws.cell(r, 5, f"=INDEX(B{r}:D{r},MATCH($C$3,$B$5:$D$5,0))").number_format = '+0.0%;-0.0%;0.0%'
        wb.defined_names[key] = DefinedName(key, attr_text=f"Inputs!$E${r}")
    r = 7 + len(DRIVERS)
    head_row(ws, r, ["Fixed input", "Value", "Source", "", ""])
    for key, (v, label, fmt, src) in X.items():
        r += 1
        ws.cell(r, 1, label)
        ws.cell(r, 2, v).number_format = fmt
        ws.cell(r, 2).font = Font(color=INPUT)
        ws.cell(r, 3, src)
        wb.defined_names[key] = DefinedName(key, attr_text=f"Inputs!$B${r}")
    r += 2
    head_row(ws, r, ["Option assumption", "Value", "Source", "", ""])
    for key, (v, label, fmt) in OPTIONS.items():
        r += 1
        ws.cell(r, 1, label)
        ws.cell(r, 2, v).number_format = fmt
        ws.cell(r, 2).font = Font(color=INPUT)
        ws.cell(r, 3, "Fictional quote or company assumption")
        wb.defined_names[key] = DefinedName(key, attr_text=f"Inputs!$B${r}")
    ws.column_dimensions["A"].width = 62
    for c in "BCDE":
        ws.column_dimensions[c].width = 12
    ws.column_dimensions["C"].width = 14

    # Actuals
    ws = wb.create_sheet("Actuals")
    header(ws, "FY26 actuals (£)", "From the P1 pack. The P2 close reproduces these figures from the ledger.")
    head_row(ws, 5, [""] + bp.MONTHS)
    RA, r = {}, 5
    for ch in CH:
        for line in ["Orders", "Gross sales", "Returns", "Cost of goods sold"]:
            r += 1
            RA[(ch, line)] = r
            ws.cell(r, 1, f"{ch} {line.lower()}")
            for p in range(12):
                ws.cell(r, 2 + p, round(A[(ch, line)][p], 2)).number_format = "#,##0"
    for line in ["Payroll", "Other overheads"]:
        r += 1
        RA[line] = r
        ws.cell(r, 1, line)
        for p in range(12):
            ws.cell(r, 2 + p, round(A[("Central", line)][p], 2)).number_format = "#,##0"
    r += 1
    RA["cogs"] = r
    ws.cell(r, 1, "Cost of goods sold, total").font = Font(bold=True)
    for p in range(12):
        L = col(2 + p)
        ws.cell(r, 2 + p, f"={L}{RA[('Website', 'Cost of goods sold')]}+{L}{RA[('Marketplace', 'Cost of goods sold')]}").number_format = "#,##0"
    ws.column_dimensions["A"].width = 34

    # Forecast
    ws = wb.create_sheet("Forecast")
    header(ws, "Rolling forecast, Feb-26 to Jul-27 (£k)", "Scenario set on the Inputs sheet.")
    ws["A3"] = '="Scenario: "&Inputs!$C$3'
    ws["A3"].font = Font(bold=True)
    head_row(ws, 5, [""] + [""] * N + ["FY27", "18 months"])
    for c, dt in enumerate(FM):
        ws.cell(5, 2 + c, dt).number_format = "mmm-yy"
    R, r = {}, 6

    for ch in CH:
        g = "web_growth" if ch == "Website" else "mkt_growth"
        k = {"Website": "web", "Marketplace": "mkt"}[ch]
        ws.cell(r, 1, ch).font = Font(bold=True, color=bp.NAVY)
        r += 1
        block = [("orders", "Orders (000s)", "#,##0,", "sum"), ("aov", "Average order value (£)", "£0.00", None),
                 ("gross", "Gross sales", GBP, "sum"), ("rr", "Returns rate", PCT, None),
                 ("returns", "Returns", GBP, "sum"), ("net", "Net revenue", GBP, "sum"),
                 ("gm", "Gross margin %", PCT, None), ("cogs", "Cost of goods sold", GBP, "sum"),
                 ("delivery", "Delivery and fulfilment", GBP, "sum"),
                 ("marketing" if ch == "Website" else "fees", "Marketing *" if ch == "Website" else "Marketplace fees", GBP, "sum"),
                 ("contribution", "Contribution", GBP, "sum")]
        for i, (name, *_) in enumerate(block):
            R[f"{k}:{name}"] = r + i
        last = "marketing" if ch == "Website" else "fees"
        for name, label, fmt, total in block:
            ws.cell(r, 1, label)
            for c in range(N):
                L, P = col(2 + c), col(2 + c % 12)
                ly = lambda line: f"Actuals!{P}{RA[(ch, line)]}"
                me = lambda n: f"{L}{R[f'{k}:{n}']}"
                f = {
                    "orders": lambda: f"={ly('Orders')}*(1+{g})" if c < 12 else f"={col(2 + c - 12)}{r}*(1+{g})",
                    "aov": lambda: f"={ly('Gross sales')}/{ly('Orders')}*(1+aov)" if c < 12 else f"={col(2 + c - 12)}{r}*(1+aov)",
                    "gross": lambda: f"={me('orders')}*{me('aov')}",
                    "rr": lambda: f"={ly('Returns')}/{ly('Gross sales')}+returns_pp",
                    "returns": lambda: f"=-{me('gross')}*{me('rr')}",
                    "net": lambda: f"={me('gross')}+{me('returns')}",
                    "gm": lambda: f"=1-{ly('Cost of goods sold')}/({ly('Gross sales')}-{ly('Returns')})+gm_pp",
                    "cogs": lambda: f"=-{me('net')}*(1-{me('gm')})",
                    "delivery": lambda: f"=-{me('orders')}*delivery_rate*(1+carrier)",
                    "marketing": lambda: f"=-{me('net')}*mkt_ratio",
                    "fees": lambda: f"=-{me('net')}*fee_rate",
                    "contribution": lambda: f"={me('net')}+{me('cogs')}+{me('delivery')}+{me(last)}",
                }[name]()
                ws.cell(r, 2 + c, f).number_format = fmt
            if total:
                ws.cell(r, 2 + N, f"=SUM(B{r}:M{r})").number_format = fmt
                ws.cell(r, 3 + N, f"=SUM(B{r}:{col(1 + N)}{r})").number_format = fmt
            style_row(ws, r, 3 + N, fmt, bold=name in ("net", "contribution"), border=name in ("net", "contribution"))
            r += 1
        r += 1

    ws.cell(r, 1, "Central").font = Font(bold=True, color=bp.NAVY)
    r += 1
    central = [
        ("orders", "Total orders (000s)", "#,##0,", lambda L, c: f"={L}{R['web:orders']}+{L}{R['mkt:orders']}", True),
        ("hours", "Warehouse hours (000s)", "#,##0,", lambda L, c: f"={L}{R['orders']}/(orders_per_hour*(1+productivity))", True),
        ("nlw", "National Living Wage (£/hour)", "£0.00", lambda L, c: f"=IF({L}$5<DATE(2026,4,1),nlw_2025,IF({L}$5<DATE(2027,4,1),nlw_2026,nlw_2026*(1+nlw_2027)))", False),
        ("warehouse", "Warehouse payroll", GBP, lambda L, c: f"=-{L}{R['hours']}*{L}{R['nlw']}*(1+oncost)", True),
        ("salaried", "Salaried payroll", GBP, lambda L, c: f"=-salaried_month*(1+pay_rise)^(({L}$5>=DATE(2026,4,1))+({L}$5>=DATE(2027,4,1)))", True),
        ("overheads", "Other overheads", GBP, lambda L, c: f"=-Actuals!{col(2 + c % 12)}{RA['Other overheads']}*(1+overheads)^{2 if c >= 12 else 1}", True),
        ("ebitda", "EBITDA", GBP, lambda L, c: f"={L}{R['web:contribution']}+{L}{R['mkt:contribution']}+{L}{R['warehouse']}+{L}{R['salaried']}+{L}{R['overheads']}", True),
        ("dep", "Depreciation", GBP, lambda L, c: "=-dep_m", True),
        ("op", "Operating profit", GBP, lambda L, c: f"={L}{R['ebitda']}+{L}{R['dep']}", True),
    ]
    for name, label, fmt, fn, total in central:
        R[name] = r
        ws.cell(r, 1, label)
        for c in range(N):
            ws.cell(r, 2 + c, fn(col(2 + c), c))
        if total:
            ws.cell(r, 2 + N, f"=SUM(B{r}:M{r})")
            ws.cell(r, 3 + N, f"=SUM(B{r}:{col(1 + N)}{r})")
        style_row(ws, r, 3 + N, fmt, bold=name in ("ebitda", "op"), border=name in ("ebitda", "op"))
        r += 1
    r += 1

    ws.cell(r, 1, "Cash").font = Font(bold=True, color=bp.NAVY)
    r += 1
    opening_trail = f"SUM(Actuals!$B${RA['cogs']}:$M${RA['cogs']})"
    cash_rows = ["cogs_total", "trail", "inv", "pay", "open", "ebitda_c", "d_inv", "d_pay", "ocf", "capex", "tax", "close"]
    for i, name in enumerate(cash_rows):
        R[name] = r + i
    labels = {"cogs_total": "Cost of goods sold, total", "trail": "Last 12 months' cost of goods sold", "inv": "Inventory",
              "pay": "Trade payables", "open": "Opening cash", "ebitda_c": "EBITDA", "d_inv": "(Increase)/decrease in inventory",
              "d_pay": "Increase/(decrease) in payables", "ocf": "Operating cash flow", "capex": "Capital expenditure",
              "tax": "Corporation tax paid", "close": "Closing cash"}
    for name in cash_rows:
        rr = R[name]
        ws.cell(rr, 1, labels[name])
        for c in range(N):
            L, prev = col(2 + c), col(1 + c)
            f = {
                "cogs_total": f"=-({L}{R['web:cogs']}+{L}{R['mkt:cogs']})",
                "trail": (f"=SUM(Actuals!{col(3 + c)}{RA['cogs']}:Actuals!$M${RA['cogs']})+SUM($B{R['cogs_total']}:{L}{R['cogs_total']})"
                          if c < 11 else f"=SUM({col(2 + c - 11)}{R['cogs_total']}:{L}{R['cogs_total']})"),
                "inv": f"=inv_days/365*{L}{R['trail']}",
                "pay": f"=pay_days/365*{L}{R['trail']}",
                "open": "=cash_open" if c == 0 else f"={prev}{R['close']}",
                "ebitda_c": f"={L}{R['ebitda']}",
                "d_inv": (f"=inv_days/365*{opening_trail}-{L}{R['inv']}" if c == 0 else f"={prev}{R['inv']}-{L}{R['inv']}"),
                "d_pay": (f"={L}{R['pay']}-pay_days/365*{opening_trail}" if c == 0 else f"={L}{R['pay']}-{prev}{R['pay']}"),
                "ocf": f"={L}{R['ebitda_c']}+{L}{R['d_inv']}+{L}{R['d_pay']}",
                "capex": "=-capex_m",
                "tax": ("=-tax_rate*fy26_profit/4" if c in (0, 3) else
                        f"=IF(SUM($B${R['op']}:$M${R['op']})>large_threshold,-tax_rate*SUM($B${R['op']}:$M${R['op']})/4,0)"
                        if c in (6, 9, 12, 15) else 0),
                "close": f"={L}{R['open']}+{L}{R['ocf']}+{L}{R['capex']}+{L}{R['tax']}",
            }[name]
            ws.cell(rr, 2 + c, f)
        if name in ("ebitda_c", "d_inv", "d_pay", "ocf", "capex", "tax"):
            ws.cell(rr, 2 + N, f"=SUM(B{rr}:M{rr})")
            ws.cell(rr, 3 + N, f"=SUM(B{rr}:{col(1 + N)}{rr})")
        elif name == "close":
            ws.cell(rr, 2 + N, f"=M{rr}")
            ws.cell(rr, 3 + N, f"={col(1 + N)}{rr}")
        elif name == "open":
            ws.cell(rr, 2 + N, f"=B{rr}")
            ws.cell(rr, 3 + N, f"=B{rr}")
        style_row(ws, rr, 3 + N, GBP, bold=name in ("ocf", "close", "open"), border=name in ("ocf", "close"))
        if name in ("cogs_total", "trail"):
            for c in range(1, 4 + N):
                ws.cell(rr, c).font = Font(italic=True, color="7F7F7F")
    r = R["close"] + 2
    ws.cell(r, 1, "* Marketing uses a placeholder assumption with no citation yet.").font = bp.F_NOTE
    ws.column_dimensions["A"].width = 34
    for c in range(2, 4 + N):
        ws.column_dimensions[col(c)].width = 9.5
    ws.freeze_panes = "B6"

    # Options
    ws = wb.create_sheet("Options")
    header(ws, "Options: cash impact against doing nothing (£k)", "Contribution only, before tax and working capital. Scenario set on the Inputs sheet.")
    ws["A3"] = '="Scenario: "&Inputs!$C$3'
    ws["A3"].font = Font(bold=True)
    head_row(ws, 5, [""] + [""] * N + ["FY27", "18 months"])
    for c, dt in enumerate(FM):
        ws.cell(5, 2 + c, dt).number_format = "mmm-yy"
    F = lambda key, L: f"Forecast!{L}{R[key]}"
    rows = [
        ("cpo_w", "Website contribution per order (£)", "£0.00",
         lambda L, c: f"=({F('web:net', L)}+{F('web:cogs', L)}+{F('web:delivery', L)})/{F('web:orders', L)}", False),
        ("cpo_m", "Marketplace contribution per order (£)", "£0.00",
         lambda L, c: f"=({F('mkt:net', L)}+{F('mkt:cogs', L)}+{F('mkt:delivery', L)}+{F('mkt:fees', L)})/{F('mkt:orders', L)}", False),
        ("A", "A. Switch carrier", GBP,
         lambda L, c: "=-switch_cost" if c == 1 else (f"=(delivery_rate*(1+carrier)-quote)*{F('orders', L)}" if c >= 3 else 0), True),
        ("B", "B. Delivery charge on small website orders", GBP,
         lambda L, c: (f"={F('web:orders', L)}*affected*((1-loss)*charge_exvat-loss*{L}{{cpo_w}})" if c >= 2 else 0), True),
        ("C", "C. 2% price rise", GBP,
         lambda L, c: (f"={F('web:orders', L)}*((1-volume_loss)*({L}{{cpo_w}}+price_rise*{F('web:net', L)}/{F('web:orders', L)})-{L}{{cpo_w}})"
                       f"+{F('mkt:orders', L)}*((1-volume_loss)*({L}{{cpo_m}}+price_rise*{F('mkt:net', L)}/{F('mkt:orders', L)}*(1-fee_rate))-{L}{{cpo_m}})"
                       if c >= 2 else 0), True),
    ]
    O = {}
    for i, (name, label, fmt, fn, total) in enumerate(rows):
        O[name] = 6 + i
    for name, label, fmt, fn, total in rows:
        rr = O[name]
        ws.cell(rr, 1, label)
        for c in range(N):
            f = fn(col(2 + c), c)
            if isinstance(f, str):
                f = f.replace("{cpo_w}", str(O["cpo_w"])).replace("{cpo_m}", str(O["cpo_m"]))
            ws.cell(rr, 2 + c, f).number_format = fmt
        if total:
            ws.cell(rr, 2 + N, f"=SUM(B{rr}:M{rr})").number_format = fmt
            ws.cell(rr, 3 + N, f"=SUM(B{rr}:{col(1 + N)}{rr})").number_format = fmt
            for c in range(1, 4 + N):
                ws.cell(rr, c).font = Font(bold=True)
    rr = O["C"] + 2
    T = lambda key: f"Forecast!{col(2 + N)}{R[key]}"
    ws.cell(rr, 1, "Break-even, option B: share of affected orders we can lose")
    ws.cell(rr, 2, f"=charge_exvat/(charge_exvat+SUMPRODUCT(B{O['cpo_w']}:M{O['cpo_w']},Forecast!B{R['web:orders']}:M{R['web:orders']})/{T('web:orders')})").number_format = "0.0%"
    ws.cell(rr + 1, 1, "Break-even, option C: share of orders we can lose")
    lift = f"price_rise*({T('web:net')}+{T('mkt:net')}*(1-fee_rate))"
    contrib = f"({T('web:net')}+{T('web:cogs')}+{T('web:delivery')}+{T('mkt:net')}+{T('mkt:cogs')}+{T('mkt:delivery')}+{T('mkt:fees')})"
    ws.cell(rr + 1, 2, f"={lift}/({contrib}+{lift})").number_format = "0.0%"
    ws.column_dimensions["A"].width = 42
    for c in range(2, 4 + N):
        ws.column_dimensions[col(c)].width = 9.5
    ws.freeze_panes = "B6"

    # Summary (values)
    ws = wb.create_sheet("Summary")
    header(ws, "Scenario summary (£k)", "FY27 runs Feb-26 to Jan-27. Values from src/forecast.py, which runs the same model as the Forecast sheet.")
    head_row(ws, 5, [""] + SCN)
    items = [
        ("FY27 net revenue", lambda F, o: sum(F["Website:net"][:12]) + sum(F["Marketplace:net"][:12]), GBP),
        ("FY27 EBITDA", lambda F, o: sum(F["ebitda"][:12]), GBP),
        ("FY27 operating profit", lambda F, o: sum(F["op"][:12]), GBP),
        ("Cash at 31 Jan 2027", lambda F, o: F["cash"][11], GBP),
        ("Lowest month-end cash, 18 months", lambda F, o: min(F["cash"]), GBP),
        ("Cash at 31 Jul 2027", lambda F, o: F["cash"][-1], GBP),
        ("Option A, 18-month cash impact", lambda F, o: sum(o["A"]), GBP),
        ("Option B, 18-month cash impact", lambda F, o: sum(o["B"]), GBP),
        ("Option C, 18-month cash impact", lambda F, o: sum(o["C"]), GBP),
    ]
    for i, (label, fn, fmt) in enumerate(items):
        ws.cell(6 + i, 1, label)
        for j, s in enumerate(SCN):
            ws.cell(6 + i, 2 + j, round(fn(*runs[s]), 2)).number_format = fmt
    ws.column_dimensions["A"].width = 40
    for c in "BCD":
        ws.column_dimensions[c].width = 12
    wb.move_sheet("Summary", offset=-4)
    wb.save(OUT)


def memo(X, runs):
    k = lambda v: f"£{abs(v) / 1e3:,.0f}k"
    sign = lambda v: ("+" if v >= 0 else "-") + k(v)
    B, o = runs["Base"]
    fy = lambda F, key: sum(F[key][:12])
    op = {s: fy(runs[s][0], "op") for s in SCN}
    low = {s: min(runs[s][0]["cash"]) for s in SCN}
    opt = {s: {x: sum(runs[s][1][x]) for x in "ABC"} for s in SCN}
    fy27 = {s: {x: sum(runs[s][1][x][:12]) for x in "ABC"} for s in SCN}
    rate = X["delivery_rate"]
    return f"""# Memo: delivery costs

> DirectSportsGoods is a fictional company and every figure here is made up. Quotes and customer responses are company assumptions, set out on the Inputs sheet of the [forecast workbook](../pack/DirectSportsGoods_forecast.xlsx).

**To:** Board  **From:** Finance  **Date:** February 2026

## The problem

Our carrier raised prices in September. Delivery now costs **£{rate:.2f} an order**, against £{bp.V['delivery']:.2f} in the budget, and the rise is permanent. At the volumes in our base forecast, that is {k(sum(B['orders'][:12]) * (rate - bp.V['delivery']))} more delivery cost in FY27 than the old rate would have meant.

The base forecast has FY27 operating profit of **{k(op['Base'])}**, against {k(X['fy26_profit'])} in FY26. The range is {k(op['Downside'])} in the downside to {k(op['Upside'])} in the upside. Cash stays positive in every scenario: the lowest month-end over the next 18 months is {k(low['Downside'])} in the downside.

## Options

Cash impact against doing nothing, before tax and working capital.

| | Base, FY27 | Base, 18 months | Downside, 18 months | Upside, 18 months |
|---|---:|---:|---:|---:|
| **A. Switch carrier** at £{OPTIONS['quote'][0]:.2f} an order, £{OPTIONS['switch_cost'][0] / 1e3:,.0f}k to switch | {sign(fy27['Base']['A'])} | {sign(opt['Base']['A'])} | {sign(opt['Downside']['A'])} | {sign(opt['Upside']['A'])} |
| **B. £1.99 delivery charge** on website orders below the free-delivery threshold | {sign(fy27['Base']['B'])} | {sign(opt['Base']['B'])} | {sign(opt['Downside']['B'])} | {sign(opt['Upside']['B'])} |
| **C. 2% price rise** across both channels | {sign(fy27['Base']['C'])} | {sign(opt['Base']['C'])} | {sign(opt['Downside']['C'])} | {sign(opt['Upside']['C'])} |

**A. Switch carrier.** The second carrier has quoted £{OPTIONS['quote'][0]:.2f} an order. Switching costs about £{OPTIONS['switch_cost'][0] / 1e3:,.0f}k and takes two months, so savings start in May. It pays back by {FM[o['payback_A']]:%B %Y}. The saving grows in the downside, because a fixed quote protects us from further rises. The risk is service: a bad peak season with a new carrier would cost more than the saving.

**B. Delivery charge.** We assume {OPTIONS['affected'][0]:.0%} of website orders fall below the threshold and {OPTIONS['loss'][0]:.0%} of those customers walk away. A lost order costs us about £{o['cpo_web']:.0f} of contribution, against £{OPTIONS['charge_exvat'][0]:.2f} from each charge, so the option only pays if fewer than **{o['be_B']:.0%}** walk away. At our assumption it roughly breaks even in every scenario, which isn't worth the risk to customer goodwill.

**C. Price rise.** The biggest prize. A 2% rise pays as long as we lose fewer than **{o['be_C']:.1%}** of orders, and we assume {OPTIONS['volume_loss'][0]:.1%}. The risk sits on marketplaces, where customers compare prices side by side and competitors can undercut us.

## Recommendation

**Switch carrier (option A) now. Test the price rise (option C) on the website before extending it. Drop the delivery charge (option B).**

- Option A is the only one that doesn't depend on how customers react, and it pays in every scenario: an estimated {k(opt['Base']['A'])} over 18 months in the base case. Ask for a peak-season service clause, with penalties, in the new contract.
- Run option C on the website for eight weeks from April, where shoppers compare prices less directly than on marketplaces. If orders fall by less than {o['be_C']:.1%}, extend it to marketplaces. At full rollout it adds up to an estimated {k(opt['Base']['C'])} more.
- Option B puts customers off for roughly no money. Leave it.

## What would change this

- If the new carrier can't guarantee peak capacity in writing, stay put and push the current carrier for a rate review instead.
- If the website price test loses more than {o['be_C']:.1%} of orders, reverse it and keep prices where they are.
- If the downside starts to play out (operating profit falls to {k(op['Downside'])}), the carrier switch matters even more: a fixed quote is worth {k(opt['Downside']['A'])} in that case.
"""


def main():
    A, fixed = calibrate()
    X = {k: v[0] for k, v in fixed.items()}
    O = {k: v[0] for k, v in OPTIONS.items()}
    runs = {}
    for j, s in enumerate(SCN):
        d = {key: vals[j] for key, _, *vals in DRIVERS}
        F = forecast(A, X, d)
        runs[s] = (F, options(F, X, O, d))
    # Self-check: base revenue is FY26 grown by the driver assumptions, and the downside never beats the base.
    fy26_net = sum(A[(c, "Gross sales")][p] - A[(c, "Returns")][p] for c in CH for p in range(12))
    base_net = sum(runs["Base"][0]["Website:net"][:12]) + sum(runs["Base"][0]["Marketplace:net"][:12])
    assert 1.0 < base_net / fy26_net < 1.2, base_net / fy26_net
    assert sum(runs["Downside"][0]["op"]) < sum(runs["Base"][0]["op"]) < sum(runs["Upside"][0]["op"])
    build_workbook(A, fixed, runs)
    MEMO.write_text(memo(X, runs), encoding="utf-8")
    print(f"Wrote {OUT.relative_to(bp.ROOT)} and {MEMO.relative_to(bp.ROOT)}")
    for s in SCN:
        F, o = runs[s]
        print(f"  {s:9s} FY27 net {(sum(F['Website:net'][:12]) + sum(F['Marketplace:net'][:12])) / 1e6:5.2f}m  "
              f"op {sum(F['op'][:12]) / 1e6:5.2f}m  low cash {min(F['cash']) / 1e6:5.2f}m  Jul-27 cash {F['cash'][-1] / 1e6:5.2f}m  "
              f"A {sum(o['A']) / 1e3:+6.0f}k  B {sum(o['B']) / 1e3:+6.0f}k  C {sum(o['C']) / 1e3:+6.0f}k  "
              f"beB {o['be_B']:.0%} beC {o['be_C']:.1%} payback {o['payback_A']}")


if __name__ == "__main__":
    main()
